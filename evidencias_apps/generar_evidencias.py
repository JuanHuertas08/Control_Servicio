"""
Genera un documento Word con evidencias de acceso a las aplicaciones del ICBF.

Lee las URL y credenciales del Excel "Listado de Aplicaciones Para Pruebas",
abre cada link con un navegador (Playwright), toma capturas antes y después del
login (cuando aplica) y arma el informe en Word.

IMPORTANTE: el script solo navega y consulta, no guarda ningún registro.

Uso:
    python generar_evidencias.py "Listado_de_Aplicaciones.xlsx"
    python generar_evidencias.py "Listado.xlsx" --manual      # pausa para revisar / loguear a mano
    python generar_evidencias.py "Listado.xlsx" --cantidad 24 --hoja "Criticidad Windows"
    python generar_evidencias.py "Listado.xlsx" --desde 25   # de la app 25 en adelante
    python generar_evidencias.py "Listado.xlsx" --solo 19,23,25   # solo esas apps
    python generar_evidencias.py "Listado.xlsx" --incluir-no-probar   # lista las "No probar" como No Aplica

Por defecto prueba todas las apps de la hoja y omite las marcadas "No probar".
"""

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import sync_playwright

VERSION = "5 (2026-10-24)"

TIMEOUT_CARGA_MS = 45_000
ESPERA_TRAS_LOGIN_MS = 6_000

NO_APLICA = "No Aplica"
OK = "OK"
FALLA = "Falla"
PARCIAL = "Revisar"

COLOR_ESTADO = {
    OK: "C6EFCE",
    FALLA: "FFC7CE",
    PARCIAL: "FFEB9C",
    NO_APLICA: "D9D9D9",
}


# --------------------------------------------------------------------------- #
# Lectura del Excel
# --------------------------------------------------------------------------- #
def limpiar(valor):
    if valor is None:
        return ""
    return str(valor).replace("\xa0", " ").strip()


def extraer_url(texto):
    """Devuelve la URL de la celda. Si hay varias, prefiere la marcada como 'Actualizada'."""
    texto = limpiar(texto)
    m = re.search(r"actualizada\s*:?\s*(https?://\S+)", texto, re.IGNORECASE)
    if m:
        return m.group(1)
    m = re.search(r"https?://\S+", texto)
    return m.group(0) if m else ""


def leer_aplicaciones(ruta_excel, hoja, cantidad, omitir_no_probar=False, desde=1):
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    ws = wb[hoja]

    # Ubica la fila de encabezados (la que tiene la columna "Url").
    fila_encabezado = None
    for fila in ws.iter_rows(min_row=1, max_row=10):
        if any(limpiar(c.value).lower() == "url" for c in fila):
            fila_encabezado = fila[0].row
            break
    if fila_encabezado is None:
        sys.exit(f"No se encontró la columna 'Url' en la hoja '{hoja}'.")

    encabezados = [limpiar(c.value).lower() for c in ws[fila_encabezado]]

    def col(*nombres):
        for n in nombres:
            if n in encabezados:
                return encabezados.index(n)
        return None

    i_app = col("server roles", "aplicación", "aplicacion")
    i_url = col("url")
    i_usr = col("usuario")
    i_pwd = col("contraseña", "contrasena")
    i_obs = col("observaciones")

    apps = []
    for fila in ws.iter_rows(min_row=fila_encabezado + 1, values_only=True):
        url = extraer_url(fila[i_url]) if i_url is not None else ""
        if not url:
            continue
        app = {
            "n": len(apps) + 1,
            "aplicacion": limpiar(fila[i_app]) if i_app is not None else "",
            "url": url,
            "usuario": limpiar(fila[i_usr]) if i_usr is not None else "",
            "password": limpiar(fila[i_pwd]) if i_pwd is not None else "",
            "observaciones": limpiar(fila[i_obs]) if i_obs is not None else "",
        }
        if omitir_no_probar and es_no_probar(app):
            continue
        apps.append(app)
    apps = apps[max(desde, 1) - 1:]
    if cantidad:
        apps = apps[:cantidad]
    return apps


def es_no_probar(app):
    textos = (app["usuario"], app["password"], app["observaciones"])
    return any("no probar" in t.lower() for t in textos)


def requiere_login(app):
    return bool(app["usuario"]) and app["usuario"].upper() != "N/A" and bool(app["password"])


# --------------------------------------------------------------------------- #
# Navegación y login
# --------------------------------------------------------------------------- #
def nombre_archivo(app, sufijo):
    base = re.sub(r"[^A-Za-z0-9]+", "_", app["aplicacion"])[:40].strip("_")
    return f"{app['n']:02d}_{base}_{sufijo}.png"


def campo_visible(page, selectores):
    for sel in selectores:
        loc = page.locator(sel)
        for i in range(loc.count()):
            el = loc.nth(i)
            try:
                if el.is_visible():
                    return el
            except PlaywrightError:
                pass
    return None


def buscar_en_frames(page, selectores):
    """Busca el campo en la página principal y en sus iframes."""
    el = campo_visible(page, selectores)
    if el:
        return el
    for frame in page.frames[1:]:
        el = campo_visible(frame, selectores)
        if el:
            return el
    return None


SEL_PASSWORD = ['input[type="password"]']
SEL_USUARIO = [
    'input[autocomplete="username"]',
    'input[name*="user" i]',
    'input[id*="user" i]',
    'input[name*="usuario" i]',
    'input[id*="usuario" i]',
    'input[name*="login" i]',
    'input[id*="login" i]',
    'input[type="email"]',
    'input[type="text"]',
    "input:not([type])",
]
SEL_BOTON = [
    'button[type="submit"]',
    'input[type="submit"]',
    'button:has-text("Ingresar")',
    'button:has-text("Iniciar")',
    'button:has-text("Entrar")',
    'button:has-text("Login")',
    'button:has-text("Acceder")',
    'a:has-text("Ingresar")',
    'input[type="button"][value*="Ingresar" i]',
    'input[type="image"]',
]


def intentar_login(page, app):
    """Login genérico: llena usuario/contraseña y envía. Devuelve (exito, detalle)."""
    # Algunas apps muestran primero un botón "Ingresar" antes del formulario.
    pwd = buscar_en_frames(page, SEL_PASSWORD)
    if pwd is None:
        boton_previo = campo_visible(page, ['a:has-text("Ingresar")', 'button:has-text("Ingresar")',
                                            'a:has-text("Iniciar sesión")'])
        if boton_previo:
            boton_previo.click()
            page.wait_for_timeout(3000)
            pwd = buscar_en_frames(page, SEL_PASSWORD)
    if pwd is None:
        return False, "No se encontró el formulario de login"

    usr = buscar_en_frames(page, SEL_USUARIO)
    if usr is None:
        return False, "No se encontró el campo de usuario"

    url_antes = page.url
    usr.fill(app["usuario"])
    pwd.fill(app["password"])

    boton = buscar_en_frames(page, SEL_BOTON)
    if boton:
        boton.click()
    else:
        pwd.press("Enter")

    try:
        page.wait_for_load_state("networkidle", timeout=ESPERA_TRAS_LOGIN_MS * 3)
    except PlaywrightTimeout:
        pass
    page.wait_for_timeout(ESPERA_TRAS_LOGIN_MS)

    sigue_login = buscar_en_frames(page, SEL_PASSWORD) is not None
    if not sigue_login or page.url != url_antes:
        return True, "Login exitoso"
    return False, "Tras enviar credenciales sigue en la pantalla de login"


def probar_aplicacion(context, app, carpeta, manual):
    res = {
        "fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "http": "",
        "titulo": "",
        "capturas": [],
        "estado": FALLA,
        "detalle": "",
    }

    if es_no_probar(app):
        res["estado"] = NO_APLICA
        res["detalle"] = "Marcado como 'No probar' en el listado"
        return res

    page = context.new_page()
    try:
        respuesta = page.goto(app["url"], timeout=TIMEOUT_CARGA_MS, wait_until="domcontentloaded")
        try:
            page.wait_for_load_state("networkidle", timeout=15_000)
        except PlaywrightTimeout:
            pass
        page.wait_for_timeout(2000)

        codigo = respuesta.status if respuesta else None
        res["http"] = str(codigo or "")
        res["titulo"] = page.title()

        captura = carpeta / nombre_archivo(app, "1_carga")
        page.screenshot(path=str(captura))
        res["capturas"].append(("Carga de la página", captura))

        if codigo and codigo >= 400:
            res["detalle"] = f"La página respondió con código HTTP {codigo}"
            return res

        if requiere_login(app):
            if manual:
                input(f"   -> Revise/inicie sesión manualmente en '{app['aplicacion']}' y pulse Enter...")
                ok, detalle = True, "Login / navegación verificados manualmente"
            else:
                ok, detalle = intentar_login(page, app)

            captura = carpeta / nombre_archivo(app, "2_login")
            page.screenshot(path=str(captura))
            res["capturas"].append(("Después del login", captura))
            res["estado"] = OK if ok else PARCIAL
            res["detalle"] = f"Página carga correctamente. {detalle}"
        else:
            if manual:
                input(f"   -> Revise la navegación en '{app['aplicacion']}' y pulse Enter...")
                captura = carpeta / nombre_archivo(app, "2_navegacion")
                page.screenshot(path=str(captura))
                res["capturas"].append(("Navegación", captura))
            res["estado"] = OK
            if app["usuario"] or "usuario" in app["observaciones"].lower():
                res["detalle"] = "Página carga correctamente. Sin credenciales en el listado para probar login"
            else:
                res["detalle"] = "Página carga correctamente (no requiere autenticación)"

    except PlaywrightTimeout:
        res["detalle"] = f"Tiempo de espera agotado ({TIMEOUT_CARGA_MS // 1000}s). ¿VPN conectada?"
    except PlaywrightError as e:
        msg = str(e).splitlines()[0]
        res["detalle"] = f"Error de acceso: {msg}"
    finally:
        page.close()
    return res


# --------------------------------------------------------------------------- #
# Documento Word
# --------------------------------------------------------------------------- #
def sombrear(celda, color_hex):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def fila_tabla(tabla, etiqueta, valor, color=None):
    celdas = tabla.add_row().cells
    celdas[0].text = etiqueta
    celdas[0].paragraphs[0].runs[0].bold = True
    sombrear(celdas[0], "DEEAF6")
    celdas[1].text = valor
    if color:
        sombrear(celdas[1], color)
        celdas[1].paragraphs[0].runs[0].bold = True


def generar_word(apps, resultados, salida, hoja, inicio, fin):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # Portada
    titulo = doc.add_heading("Evidencias de Pruebas de Acceso a Aplicaciones", 0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Hoja: {hoja}  |  Aplicaciones: {len(apps)}\n").bold = True
    p.add_run(f"Fecha de ejecución: {inicio:%Y-%m-%d %H:%M} a {fin:%H:%M}")
    aviso = doc.add_paragraph()
    aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = aviso.add_run("Prueba de solo navegación y consulta. No se guardó ningún registro.")
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # Resumen
    conteo = {k: 0 for k in (OK, PARCIAL, FALLA, NO_APLICA)}
    for res in resultados:
        conteo[res["estado"]] += 1
    doc.add_heading("Resumen", level=1)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, estado in zip(t.rows[0].cells, conteo):
        celda.text = f"{estado}: {conteo[estado]}"
        celda.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        celda.paragraphs[0].runs[0].bold = True
        sombrear(celda, COLOR_ESTADO[estado])

    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for celda, txt in zip(t.rows[0].cells, ["#", "Aplicación", "URL", "Estado", "Observación"]):
        celda.text = txt
        celda.paragraphs[0].runs[0].bold = True
        sombrear(celda, "1F4E79")
        celda.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for app, res in zip(apps, resultados):
        c = t.add_row().cells
        c[0].text = str(app["n"])
        c[1].text = app["aplicacion"]
        c[2].text = app["url"]
        c[3].text = res["estado"]
        sombrear(c[3], COLOR_ESTADO[res["estado"]])
        c[4].text = res["detalle"]
        for celda in c:
            for par in celda.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8)
    anchos = [Cm(0.8), Cm(4.5), Cm(5.2), Cm(1.8), Cm(4.7)]
    for fila in t.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = ancho

    # Detalle por aplicación
    for app, res in zip(apps, resultados):
        doc.add_page_break()
        doc.add_heading(f"{app['n']}. {app['aplicacion']}", level=1)
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        fila_tabla(t, "URL", app["url"])
        fila_tabla(t, "Fecha / hora", res["fecha"])
        fila_tabla(t, "Requiere login", "Sí" if requiere_login(app) else "No")
        if requiere_login(app):
            fila_tabla(t, "Usuario de prueba", app["usuario"])
        fila_tabla(t, "Instrucción del listado", app["observaciones"] or "-")
        fila_tabla(t, "Código HTTP", res["http"] or "-")
        fila_tabla(t, "Título de la página", res["titulo"] or "-")
        fila_tabla(t, "Resultado", res["estado"], COLOR_ESTADO[res["estado"]])
        fila_tabla(t, "Observación", res["detalle"])
        for fila in t.rows:
            fila.cells[0].width = Cm(4.5)
            fila.cells[1].width = Cm(12.5)

        for leyenda, ruta in res["capturas"]:
            doc.add_paragraph()
            doc.add_paragraph(leyenda).runs[0].bold = True
            doc.add_picture(str(ruta), width=Cm(16.5))

    doc.save(salida)


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description="Genera evidencias de acceso a aplicaciones en Word.")
    ap.add_argument("excel", nargs="?", default="Listado_de_Aplicaciones_Para_Pruebas_de_Aplicaciones_2026.xlsx",
                    help="Ruta del Excel con el listado de aplicaciones")
    ap.add_argument("--hoja", default="Criticidad Windows")
    ap.add_argument("--cantidad", type=int, default=0, help="Número de links a probar (desde el primero); 0 = todos")
    ap.add_argument("--desde", type=int, default=1, help="Número de la app desde la que se empieza (ej. 25)")
    ap.add_argument("--solo", default="", help="Números de las apps a probar, separados por coma (ej. 19,23,25)")
    ap.add_argument("--incluir-no-probar", action="store_true",
                    help="Incluye las apps marcadas 'No probar' en el informe como 'No Aplica'")
    ap.add_argument("--salida", default=None, help="Nombre del archivo Word de salida")
    ap.add_argument("--manual", action="store_true",
                    help="Pausa en cada app para revisar / iniciar sesión manualmente")
    ap.add_argument("--oculto", action="store_true", help="Ejecuta el navegador sin ventana")
    args = ap.parse_args()

    apps = leer_aplicaciones(args.excel, args.hoja, args.cantidad, not args.incluir_no_probar, args.desde)
    if args.solo:
        numeros = {int(n) for n in re.findall(r"\d+", args.solo)}
        apps = [a for a in apps if a["n"] in numeros]
    inicio = datetime.now()
    carpeta = Path(f"evidencias_{inicio:%Y%m%d_%H%M}")
    carpeta.mkdir(exist_ok=True)
    salida = args.salida or str(carpeta / f"Evidencias_Aplicaciones_{inicio:%Y%m%d_%H%M}.docx")

    print(f"generar_evidencias.py - versión {VERSION}")
    print(f"Excel: {Path(args.excel).resolve()}")
    print(f"Probando {len(apps)} aplicaciones de la hoja '{args.hoja}'...\n")
    resultados = []
    with sync_playwright() as pw:
        navegador = pw.chromium.launch(headless=args.oculto and not args.manual)
        context = navegador.new_context(ignore_https_errors=True, viewport={"width": 1366, "height": 768})
        for app in apps:
            print(f"[{app['n']:02d}/{len(apps)}] {app['aplicacion']} - {app['url']}")
            res = probar_aplicacion(context, app, carpeta, args.manual)
            print(f"        {res['estado']}: {res['detalle']}")
            resultados.append(res)
            # Cierra sesión entre apps para que el SSO no afecte la siguiente prueba.
            context.clear_cookies()
        navegador.close()

    generar_word(apps, resultados, salida, args.hoja, inicio, datetime.now())
    print(f"\nInforme generado: {salida}")


if __name__ == "__main__":
    main()
